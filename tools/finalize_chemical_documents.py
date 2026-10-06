"""Create a reviewable tracker copy and source-backed commercial content artifacts."""
from pathlib import Path
from copy import copy
from lxml import html
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
import csv,json

ROOT=Path(__file__).resolve().parents[1]
DOCS=ROOT/'docs/chemical-build'
WORK=DOCS/'content'
SITE='https://internationalprocessplants.com'
INPUT=Path.home()/'Downloads/Used Chemical Process Plants for Sale _ International Process Plants .xlsx'

dispositions={
 2:('Draft implemented','All page contact CTAs use /contact/. Browser verified the live form. Translated-page and missing PL/HI links require CMS updates.'),
 3:('Draft implemented','Chemical Engineering replacement loads. Browser verification found the old AIChE URL is a rendered 404, so it is omitted.'),
 4:('Draft implemented','Plain unlinked H2 headings; Why heading uses the revised Suggested Outline wording. Original seven categories retained.'),
 5:('Draft implemented','CollectionPage + Organization + ItemList + BreadcrumbList + FAQPage. No Article or Person node. Apply equivalent Yoast configuration in CMS.'),
 6:('Draft implemented','og:type website; no article:published_time. Static draft does not modify WordPress filters.'),
 7:('Draft implemented','No author email, reading time or share box in draft. Sitewide WordPress user display-name cleanup is still external to this repository.'),
 8:('Draft implemented','Six relevant chemical/plant resources at the bottom; latest-posts feed removed.'),
 9:('Draft implemented','Real IMS photographs, descriptive alt text, WebP photographs under 200 KB each. No boiler image or invented stock photo. Missing plant photograph is stated honestly.'),
 10:('Draft implemented','All seven original descriptions preserved; specimen-specific materials added with verified hub and IMS links.'),
 11:('Draft implemented; confirm selection','Eight manufacturers from real featured IMS records, each with a validated manufacturer filter URL. Selection is representative, not ranked as leading. IPP can confirm which makers to feature.'),
 12:('Draft implemented','Buyer steps cover specifications, inspection, purchase, dismantling, relocation, shipping, start-up scope, financing, GPS and UGE.'),
 13:('Draft implemented','Buyer-first seller section covers asset types, documentation, valuation, decommissioning, dismantling and marketing. Seller FAQs and both official submission destinations included. No redirects applied.'),
 14:('Draft implemented','Three published warehouses, global support, asset-level locations and date on company-wide footprint. No location filter. Correction: current IMS plant listings DO expose location when supplied.'),
 15:('Draft implemented','BASF, Chemical Processing, ISN announcement, CDMO, GEO/SOCMA, Novartis and Ross Gale article linked. ISN described as a 2024 announcement, not independently certified current status. Company figures dated.'),
 16:('Needs real reviewer sign-off','Ross Gale’s published industry perspective linked. No claim that he reviewed this new page and no reviewedBy schema until an actual review is confirmed.'),
 17:('Draft implemented','Three key original answers preserved. Concise linked definition plus buyer/seller FAQs for documentation, financing, shipping, assessment schedule and site outcome.'),
 18:('Draft implemented; founding year unresolved','Full brand first mention, consistent HQ/phone, About link and og:site_name. Founding year and exact legal-company naming need authoritative confirmation; conflicting dates are not inserted.'),
 19:('Draft implemented with preservation exception','No em dash, defect-led content, comparison advice or invented sectors. Existing dated company-wide figures retained at user request; live stock counts are not claimed. Skill banned-word hit for “addition” is the preserved original FAQ phrase.'),
 20:('CMS action pending','No redirect created. Manual Inputs and Tracker conflict over seller-page destinations. Resolve a canonical map including translations before CMS changes.'),
 21:('CMS action pending','Same redirect-map conflict. Page copy is prepared; no source post was changed or unpublished.'),
 22:('CMS action pending','Same redirect-map conflict. Seller destination itself differs between sheet instructions. Requires one approved destination map.'),
 23:('CMS action pending','Destination explicitly unresolved in workbook: this page or /plants/. No redirect created.'),
 24:('Publication action pending','Draft retains production canonical and is noindex,nofollow. Live deployment, redirect checks, Search Console submission and monitoring were not performed.'),
 25:('Sitewide CMS action pending','Repository is a static staging project, not the WordPress navigation theme. No live Industries menu/homepage mutation.'),
 26:('Sitewide CMS action pending','Outbound equipment-hub links are in the draft. Inbound “Browse by industry: Chemical” links must be added to production hubs.'),
 27:('Consolidation review pending','Broader chemical-post consolidation needs page-by-page preservation and an agreed redirect map. No production consolidation performed.'),
 28:('External directory action pending','No ThomasNet account/listing update performed.'),
 29:('Outreach not sent','No publisher or partner messages sent. Workbook suggestions are not authorization to message third parties.'),
}
book=load_workbook(INPUT)
sheet=book['Tracker']
sheet.cell(1,5,'Draft disposition')
sheet.cell(1,6,'Implementation evidence / remaining work')
for row,(status,note) in dispositions.items():
 sheet.cell(row,5,status);sheet.cell(row,6,note)
 for col in [5,6]:
  cell=sheet.cell(row,col);cell.alignment=Alignment(vertical='top',wrap_text=True);cell.font=Font(name='Calibri',size=11,color='183947')
 sheet.row_dimensions[row].height=max(sheet.row_dimensions[row].height or 15,78)
for c in sheet[1][:6]: c.fill=PatternFill('solid',fgColor='183947');c.font=Font(name='Calibri',size=11,bold=True,color='FFFFFF');c.alignment=Alignment(wrap_text=True,vertical='center')
sheet.column_dimensions['E'].width=35;sheet.column_dimensions['F'].width=105;sheet.freeze_panes='B2';sheet.auto_filter.ref='A1:F29'
notes=book.create_sheet('Build Review')
for row in [
 ['Optimization draft','October 7, 2026'],
 ['Canonical URL',SITE+'/chemical-process-plant-equipment-used-systems-for-sale/'],
 ['Scope','Full optimization draft in one build; original source workbook remains untouched.'],
 ['Branch','codex/chemical-process-plants'],
 ['Existing value','Retained category names/order/descriptions, three principal FAQ answers, company description and original sectors.'],
 ['Revised outline','Broader plants-for-sale H1 follows the user’s updated outline. The earlier equipment-only recommendation is superseded.'],
 ['Assets','Real IPP logo/customer marks; Montserrat, Roboto, Inter; original IMS photographs encoded to WebP.'],
 ['Inventory','Six plants and eight equipment records checked against IMS. Snapshot, not an automatic live inventory feed.'],
 ['Publication','Local draft only; no push, deployment, CMS update, redirect or outreach.'],
 ['Reviewer','Ross Gale is a linked published contributor, not a claimed reviewer of this draft.'],
 ['Source exception','Old AIChE URL confirmed as an actual 404 in browser and omitted.'],
 ['Research','DataForSEO failed with an upstream internal search-engine error after location correction. No AI Overview or search-volume evidence is claimed; public primary-source verification was used.'],
 ['Original Status column','Preserved to avoid implying the live site has been changed. New columns track local draft work.'],
]: notes.append(row)
notes.column_dimensions['A'].width=28;notes.column_dimensions['B'].width=115
for row in notes:
 for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
 notes.row_dimensions[row[0].row].height=42
book.save(DOCS/'IPP-Chemical-Optimization-Tracker.xlsx')

outline=json.loads((WORK/'outline.json').read_text(encoding='utf-8'))
records=json.loads((DOCS/'display-inventory.json').read_text(encoding='utf-8'))
verified={
 'Brand':{'values':['International Process Plants (IPP)'],'source':SITE+'/'},
 'Headquarters':{'values':['410 Princeton Hightstown Road, Princeton Junction, NJ 08550, USA','+1 609-586-8004'],'source':SITE+'/contact/'},
 'Company-wide footprint':{'values':['15,000+ pieces of inventory','20 complete process plant sites','15 countries','Checked October 7, 2026; retained existing company description'],'source':SITE+'/'},
 'Financing':{'values':['Financing Available; terms subject to discussion with IPP'],'source':SITE+'/'},
 'Industries we serve':{'values':['Petrochemicals','Fertilizers','Polymers and plastics','Solvents, resins and adhesives','Acids, amines and surfactants','Agrochemicals and specialty chemicals'],'source':SITE+'/chemical-process-plant-equipment-used-systems-for-sale/'},
 'Warehouses':{'values':['Eastover, SC, USA','Billingham, UK','Bitterfeld-Wolfen, Germany'],'source':SITE+'/buy-used/'},
}
fact_records=[];eav=[]
for rec in records:
 vals=dict(rec['display_specs'])
 for k in ['manufacturer','condition','material']:
  if rec.get(k):vals[k]=rec[k]
 verified['IPP stock '+rec['id']]={'values':[f'{k}: {v}' for k,v in vals.items()],'source':rec['url']}
 for attr,value in vals.items():
  eav.append(['Used chemical process plants',rec['display_title'],attr,value,rec['url'],'Inventory card: '+rec['id']])
  fact_records.append({'heading':outline[3]['topic'],'attribute':rec['id']+' / '+attr,'value':value,'unit':'As published or equivalent conversion','source':rec['url'],'authority_tier':'tier4_self','confidence':1.0})
for k,rec in verified.items():
 if k.startswith('IPP stock'):continue
 for v in rec['values']:eav.append(['Used chemical process plants','International Process Plants',k,v,rec['source'],'Company / contact / industries / locations'])
with (DOCS/'entity-associated-entities-closed-attributes.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.writer(f);w.writerow(['Primary entity','Associated entity','Attribute','Closed value','Source','Owning section']);w.writerows(eav)
job=json.loads((WORK/'job.json').read_text(encoding='utf-8'))
job['verified_facts']=verified
job['seo_meta']={'seo_title':'Used Chemical Process Plants for Sale | IPP','meta_description':'Explore used chemical process plants, complete lines and equipment from IPP. View specifications, arrange inspections, or sell your surplus chemical assets.','h1':'Used Chemical Process Plants for Sale'}
# Explicit target keeps the skill exporter from inventing a new keyword-derived URL.
job['topical_map']=[{'keyword':'used chemical process plants for sale','title':'Used Chemical Process Plants for Sale','slug':'/chemical-process-plant-equipment-used-systems-for-sale/'}]
(WORK/'job.json').write_text(json.dumps(job,indent=2,ensure_ascii=False),encoding='utf-8')
(WORK/'facts.json').write_text(json.dumps(fact_records,indent=2,ensure_ascii=False),encoding='utf-8')
notes=json.loads((WORK/'review_notes.json').read_text(encoding='utf-8'))
notes += [
 'RESOLVED UNHOSTED_FACT: Industries we serve is explicitly present as the six original sectors in H009; the post-pass matcher failed to associate it.',
 'RESOLVED NO_QUESTION_FORMAT: headings follow the user-authored Suggested Outline; FAQs contain twelve buyer/seller questions. No quota-driven heading changes.',
 'SCOPED STYLE EXCEPTION: the only remaining banned word, addition, is in the exact retained FAQ answer “Yes. In addition to individual equipment, we sell ...”. User preservation requirement overrides that lexical ban. No new use of the banned term is accepted.',
 'DataForSEO initial location string rejected; corrected request returned Internal SE Server Error (40101). No AI Overview, PAA consensus or search-volume claim used. Public first-party IPP/IMS verification governs all factual copy.',
 'Old AIChE URL checked in browser: “Oh oh! 404”; removed. Contact form loads in browser despite automated HTTP403.',
 'Hero finalized from completed sections: complete plants, equipment, selling; preserves opening equipment entity and original readiness value.',
]
notes=[{'phase':'editorial_review','notes':[n]} if isinstance(n,str) else n for n in notes]
(WORK/'review_notes.json').write_text(json.dumps(notes,indent=2,ensure_ascii=False),encoding='utf-8')
seg={'context_segments':[{'context':n['topic'],'representative_query':n.get('optimized_query_network',''),'intent':'Commercial','commercial_layer':'Buyer-led collection; supporting seller intent','search_journey_position':str(n['funnel_stage']),'bucket':n.get('section_class',''),'total_sv':'','kw_count':'','queries':[n.get('optimized_query_network')] if n.get('optimized_query_network') else []} for n in outline],'query_dimensions':{'dimensions':[]},'vector_contexts':[],'page_signals':[],'optimized_query_network':[{'keyword':q,'search_volume':'','source':'User request / original page','intent_segment':'Commercial'} for q in ['used chemical process plants for sale','used chemical plant equipment','chemical process equipment for sale']],'excluded_keywords':[]}
(WORK/'segmentation.json').write_text(json.dumps(seg,indent=2,ensure_ascii=False),encoding='utf-8')
print('Created tracker copy, factual entity map and content export inputs.')
